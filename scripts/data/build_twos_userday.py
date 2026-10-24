"""E-11: build the TWOS user-day table (logon-only mapping) used as the fifth member (C5).

TWOS is released by its authors under an agreement and must not be redistributed. Request it from them
(Harilal et al., 2017); this script only builds the table from your own copy.

Inputs (from the TWOS release, in --twos-dir):
  eventviewer_ano.zip   Windows event logs per machine
  ImportantInfo.xlsx    machine-to-owner map and the masquerade sessions
Features: user_day_features(deviation=True) from the cross-dataset study repository (--ijacsa-repo).

Event mapping (logon events only; servers and unmapped machines are dropped):
  LogInSuccess -> logon, LogOff -> logoff, LogInAttempt / LogInFailed -> auth.
  Machine -> owner from ImportantInfo.xlsx (rows 2-25: column A = user, column G = machine, column I = second machine).
Label rule (user-day level):
  * masquerade victims: the user named in ImportantInfo.xlsx column M, rows 3-8 on Tue 21 March 2017 and
    rows 13-18 on Thu 23 March 2017;
  * traitors: each traitor user's own user-day on Wed 22 March 2017. The five traitor user names are given with
    --traitors, as listed in the TWOS release documentation.
Outputs (in --out):
  twos_all_days.parquet               all days
  twos_user_day_dev_v1_logon.parquet  competition week only (20-24 March 2017); the table the experiments read.
                                      run_r2.py checks its SHA-256 before use. The reference table was written
                                      with pyarrow 25.0.1; another pyarrow version can change the file bytes (and
                                      so the hash) without changing the table.

Run:  python build_twos_userday.py --ijacsa-repo <clone> --twos-dir <TWOS release folder> --traitors U1 U2 U3 U4 U5
"""
import argparse
import csv
import hashlib
import io
import os
import sys
import zipfile

import openpyxl
import pandas as pd

ap = argparse.ArgumentParser(description="Build the TWOS user-day table (logon-only).")
ap.add_argument("--ijacsa-repo", required=True, help="local clone of the cross-dataset study repository (provides src.data)")
ap.add_argument("--twos-dir", required=True, help="folder holding eventviewer_ano.zip and ImportantInfo.xlsx")
ap.add_argument("--traitors", nargs=5, required=True, metavar="USER",
                help="the five traitor users of Wed 22 March 2017, as named in the TWOS release")
ap.add_argument("--out", default=os.path.join(os.environ.get("FITD_ROOT", "."), "data", "twos"),
                help="output folder (default: $FITD_ROOT/data/twos)")
A = ap.parse_args()

sys.path.insert(0, A.ijacsa_repo)
from src.data.features import user_day_features  # noqa: E402

TW = A.twos_dir
os.makedirs(A.out, exist_ok=True)

ws = openpyxl.load_workbook(os.path.join(TW, "ImportantInfo.xlsx")).active
m2u = {}
for r in ws.iter_rows(min_row=2, max_row=25, values_only=True):
    m2u[r[6]] = r[0]
    if r[8]:
        m2u[r[8]] = r[0]

AMAP = {"LogInSuccess": "logon", "LogOff": "logoff", "LogInAttempt": "auth", "LogInFailed": "auth"}
rows = []
Z = zipfile.ZipFile(os.path.join(TW, "eventviewer_ano.zip"))
for n in Z.namelist():
    m = n.split("/")[-1][:-4]
    if m not in m2u:
        continue
    for r in csv.reader(io.TextIOWrapper(Z.open(n), "utf-8", "replace")):
        if len(r) >= 3 and r[2] in AMAP:
            rows.append((r[0], m2u[m], m, AMAP[r[2]]))
ev = pd.DataFrame(rows, columns=["ts", "user", "dst_host", "action"])
ev["timestamp"] = pd.to_datetime(ev.ts, format="%m/%d/%y %H:%M:%S")
ev["label"] = 0

# labels: user-day level
pos = set()
# masquerade victims: Tue 21 rows 3-8, Thu 23 rows 13-18 (column M)
for rr, day in ((range(3, 9), "2017-03-21"), (range(13, 19), "2017-03-23")):
    for i in rr:
        pos.add((ws.cell(i, 13).value, day))
# traitors: own user-day on Wed 22
for u in A.traitors:
    pos.add((u, "2017-03-22"))
print("positive user-days defined:", len(pos))

f = user_day_features(ev[["timestamp", "user", "dst_host", "action", "label"]], deviation=True)
f["day"] = f["day"].dt.strftime("%Y-%m-%d")
f["label"] = [int((u, d) in pos) for u, d in zip(f.user, f.day)]
missing = [p for p in pos if not ((f.user == p[0]) & (f.day == p[1])).any()]
print("positives without any logon row:", len(missing))
f.to_parquet(os.path.join(A.out, "twos_all_days.parquet"))
print(f.groupby("day").agg(rows=("label", "size"), pos=("label", "sum")).to_string())

c = f[f.day.between("2017-03-20", "2017-03-24")]
print("competition rows", len(c), "pos", c.label.sum())
p = os.path.join(A.out, "twos_user_day_dev_v1_logon.parquet")
c.to_parquet(p, index=False)
print("SHA-256", hashlib.sha256(open(p, "rb").read()).hexdigest())
